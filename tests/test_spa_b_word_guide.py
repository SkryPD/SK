from pathlib import Path
import sys
import unittest

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_spa_b_word_guide import build_content_model, build_document


class ContentModelTests(unittest.TestCase):
    def test_content_model_covers_approved_scope(self):
        model = build_content_model()
        self.assertEqual(model["project"]["location"], "广东省清远市英德市黄花镇")
        self.assertEqual(model["project"]["scheme"], "汤泉 B 方案")
        self.assertEqual(model["project"]["lighting"], "约 4000K 中性白光")
        self.assertEqual(len(model["tiles"]), 8)
        self.assertEqual(
            {row["zone"] for row in model["tiles"]},
            {
                "池底",
                "池壁",
                "坐台/入池",
                "池沿",
                "池外汤区地面",
                "过渡区地面",
                "普通墙面",
                "端景肌理墙",
            },
        )
        self.assertEqual(model["ceilings"][0]["system"], "铝蜂窝板系统")
        self.assertEqual(model["ceilings"][1]["system"], "工程铝扣板系统")
        self.assertIn("钢化夹层安全玻璃", model["glass"]["glass_type"])
        self.assertIn("磨砂防撞带", model["glass"]["visibility"])
        self.assertEqual(model["lighting"][0]["cct"], "4000K")
        self.assertTrue(all(row["cct"] == "4000K" for row in model["lighting"]))


    def test_content_model_keeps_professional_boundaries(self):
        model = build_content_model()
        joined = "\n".join(model["boundaries"])
        for term in ["结构", "防水", "电气", "消防", "暖通", "给排水", "机电"]:
            self.assertIn(term, joined)
        self.assertIn("不从效果图反推", joined)
        self.assertIn("玻璃厚度", joined)
        self.assertIn("专业单位", joined)


class DocumentBuildTests(unittest.TestCase):
    def test_build_document_has_required_sections_and_tables(self):
        output = ROOT / "tmp" / "test-spa-b-word" / "guide.docx"
        output.parent.mkdir(parents=True, exist_ok=True)
        build_document(output)
        self.assertTrue(output.exists())
        self.assertGreater(output.stat().st_size, 0)

        doc = Document(output)
        text = "\n".join(p.text for p in doc.paragraphs)
        table_text = "\n".join(
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        )
        all_text = f"{text}\n{table_text}"
        for heading in [
            "瓷砖选择",
            "天花系统",
            "玻璃隔断",
            "灯具选择指南",
            "施工验收检查表",
            "专业边界",
        ]:
            self.assertIn(heading, text)
        for term in [
            "150×150",
            "600×1200",
            "铝蜂窝板",
            "铝扣板",
            "钢化夹层",
            "磨砂防撞带",
            "4000K",
            "IP65",
            "IP67",
        ]:
            self.assertIn(term, all_text)
        self.assertGreaterEqual(len(doc.tables), 10)
        self.assertGreaterEqual(len(doc.sections), 1)
        hard_page_breaks = doc._element.xpath(".//w:br[@w:type='page']")
        self.assertEqual(
            len(hard_page_breaks),
            1,
            "Only the cover should force a hard page break; body sections must flow naturally.",
        )
        title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
        self.assertIsNone(
            title_ppr.find(qn("w:pBdr")),
            "The cover title must not inherit Word's default blue border.",
        )

        numbered_starts = [
            "机电提供新风",
            "核对样品标签",
        ]
        num_ids = []
        for start in numbered_starts:
            paragraph = next(p for p in doc.paragraphs if p.text.startswith(start))
            num_ids.append(int(paragraph._p.pPr.numPr.numId.val))
        self.assertEqual(len(set(num_ids)), len(num_ids))
        numbering = doc.part.numbering_part.element
        for num_id in num_ids:
            abstract_id = int(
                numbering.xpath(
                    f"./w:num[@w:numId='{num_id}']/w:abstractNumId"
                )[0].get(qn("w:val"))
            )
            start = numbering.xpath(
                f"./w:abstractNum[@w:abstractNumId='{abstract_id}']/w:lvl[@w:ilvl='0']/w:start"
            )[0]
            self.assertEqual(start.get(qn("w:val")), "1")


if __name__ == "__main__":
    unittest.main()
