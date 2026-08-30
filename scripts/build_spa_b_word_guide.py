"""Build the approved spa B material and construction guide.

The content model is intentionally independent from Word generation so approved
requirements can be tested without opening or rendering a document.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = (
    ROOT
    / "outputs"
    / "2026-08-30-spa-b-word"
    / "听风径民宿汤泉B方案选材与构造指南.docx"
)

# Selected preset: compact_reference_guide.
# Named project overrides: A4, 18/17 mm margins, 9.5 pt body, Chinese fonts,
# pale cool-gray/gray-green palette, and A4 content width of 9,980 DXA.
CONTENT_WIDTH_DXA = 9980
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
FONT_NAME = "微软雅黑"
FONT_FALLBACK = "Microsoft YaHei"
INK = "263331"
MUTED = "64706D"
ACCENT = "5B756E"
ACCENT_DARK = "3E5B54"
PALE_GREEN = "E8F0ED"
PALE_GRAY = "F2F4F3"
LIGHT_GRAY = "E2E7E5"
WHITE = "FFFFFF"
CAUTION = "7B6326"
PALE_CAUTION = "F7F2E5"
RISK = "8B3D3D"
PALE_RISK = "F8EAEA"


INQUIRY_FIELDS = [
    "品牌",
    "系列",
    "型号",
    "规格",
    "色号",
    "生产批次",
    "性能检测资料",
    "不含税单价",
    "含税单价",
    "运输费",
    "安装费",
    "建议损耗率",
    "交货周期",
    "质保",
    "补货周期",
    "售后联系人",
]

SAMPLE_CHECKLIST = [
    "核对样品标签与报价型号、色号、规格完全一致。",
    "瓷砖和金属面板同时在现场自然光及约 4000K 灯光下检查，不接受奶油黄、暖米黄、粉灰或橙红木色偏色。",
    "瓷砖至少组合 4 块观察整体色差、纹理频率、哑光程度和缝色；池沿制作 1:1 圆角或倒圆节点样板。",
    "筒灯、灯带和洗墙灯在真实面材附近试亮，检查色温一致性、眩光、水面倒影和端景光斑。",
    "玻璃确认超白透明度、磨砂防撞带样式、框料颜色、密封胶颜色和实体门收口关系。",
    "样板经业主、设计、施工和供应商共同签认后封样；未封样不得批量订货。",
]

ACCEPTANCE_STAGES = {
    "材料进场": [
        "核对品牌、系列、型号、规格、数量和送审封样。",
        "瓷砖锁定色号、尺码和生产批次，抽查外观、平整度及破损。",
        "查验湿态防滑、吸水率、燃烧性能、IP 等资料是否对应实际到货型号。",
        "玻璃原片、钢化和夹层标识齐全；框料、五金、密封材料有合格资料。",
    ],
    "基层/隐蔽": [
        "防水、基层强度、平整度、排水坡向和闭水试验由相应专业验收并形成记录。",
        "天花吊杆、龙骨、连接件及防腐处理完成，机电点位和检修范围已综合。",
        "玻璃顶部、侧边和底部固定基体满足深化图要求；防水和挡水收口连续。",
        "灯具回路、线管、接地、驱动位置及湿区安全措施经电气专业验收。",
    ],
    "安装过程": [
        "瓷砖按批准排版施工，池壁、池底、池沿和设备口关系清楚；空鼓、崩边及时处理。",
        "天花分格顺直，灯具、风口、喷淋、检修口位置协调；不得随意切断主龙骨。",
        "玻璃安装采用相容垫块和密封材料，固定可靠，边部受保护，实体门开启顺畅。",
        "灯具色温保持 4000K 一致，灯带接头和临水段密封按深化要求完成。",
    ],
    "完工": [
        "完成湿区防滑、空鼓、坡向、排水、收口和清洁维护检查。",
        "天花平整、无锈蚀和明显色差；全部检修口可开启，设备可维护。",
        "玻璃磨砂防撞带连续清晰，无崩边划伤；门扇净空、开启和逃生关系经复核。",
        "完成照度、眩光、回路、绝缘、防护和样灯效果验收；交付灯具及驱动清单。",
    ],
}

FORBIDDEN_ITEMS = [
    "普通石膏板",
    "天然木皮",
    "吸水矿棉板",
    "高反光湿区砖",
    "仿古黄石或吸水文化石",
    "奶油黄、暖米黄、粉灰和明显金色纹理",
    "橙红木色或厚重黑框",
    "3000K/3500K 暖光混入 4000K 照明系统",
]


def build_content_model() -> dict:
    """Return the approved, testable guide content model."""

    return {
        "project": {
            "name": "听风径民宿汤泉室内装修",
            "location": "广东省清远市英德市黄花镇",
            "scheme": "汤泉 B 方案",
            "use": "施工交底及供应商询价",
            "environment": "男女汤及两个休息/过渡区，完整封闭室内高湿环境",
            "style": "明快现代自然",
            "lighting": "约 4000K 中性白光",
        },
        "tiles": [
            {
                "zone": "池底",
                "size": "约 150×150 mm",
                "selection": "浅冷灰哑光小规格瓷质防滑砖",
                "performance": "建议吸水率≤0.5%；提交对应型号湿态防滑检测资料",
                "execution": "先排版试铺；坡向与排水由防水/给排水专业确认",
            },
            {
                "zone": "池壁",
                "size": "约 300×600 mm",
                "selection": "浅灰米白、低纹理、哑光瓷质砖",
                "performance": "建议吸水率≤0.5%；耐污、易清洁",
                "execution": "与池底和池沿对缝；阴阳角及设备口预排",
            },
            {
                "zone": "坐台/入池",
                "size": "约 150×150 mm",
                "selection": "与池底协调的浅冷灰哑光防滑瓷质砖",
                "performance": "提交对应型号湿态防滑检测资料",
                "execution": "踏步前缘圆滑、防绊；尺寸以池体深化为准",
            },
            {
                "zone": "池沿",
                "size": "约 300×600 mm 加工砖",
                "selection": "浅冷灰米白低纹理瓷质砖",
                "performance": "低吸水、耐温湿变化；加工后无崩边",
                "execution": "外露边圆角或倒圆；先做 1:1 节点样板",
            },
            {
                "zone": "池外汤区地面",
                "size": "约 600×600 mm",
                "selection": "浅冷灰米白细石纹哑光防滑砖",
                "performance": "提交对应型号湿态防滑检测资料",
                "execution": "分区找坡；避免高反光和深色拼花",
            },
            {
                "zone": "过渡区地面",
                "size": "约 600×600 mm",
                "selection": "浅冷灰米白细石纹防滑砖",
                "performance": "耐污、易清洁；湿态防滑资料与型号一致",
                "execution": "与汤区门槛、玻璃基座和挡水收口协调",
            },
            {
                "zone": "普通墙面",
                "size": "约 600×1200 mm",
                "selection": "中性灰白低纹理哑光瓷质砖",
                "performance": "低反光、耐污、易清洁",
                "execution": "贴至天花；排版扣除门、玻璃及洞口",
            },
            {
                "zone": "端景肌理墙",
                "size": "按深化排版",
                "selection": "浅灰米白浅浮雕/肌理砖",
                "performance": "低反光、耐湿、易清洁",
                "execution": "配合 4000K 均匀洗墙；不设水帘",
            },
        ],
        "ceilings": [
            {
                "tier": "中档主选",
                "system": "铝蜂窝板系统",
                "module": "约 600×1200 mm",
                "finish": "浅灰白哑光",
                "requirements": "耐潮耐腐蚀；龙骨、吊杆、螺钉和连接件同步适应凝露环境",
            },
            {
                "tier": "经济备选",
                "system": "工程铝扣板系统",
                "module": "约 600×600 mm",
                "finish": "浅灰白哑光",
                "requirements": "耐潮耐腐蚀；便于拆装检修和后续补板",
            },
        ],
        "glass": {
            "position": "休息/过渡区与汤池区之间，通高玻璃框住防潮实体门",
            "glass_type": "优先采用超白透明钢化夹层安全玻璃体系",
            "visibility": "中部设置连续磨砂防撞带",
            "frame": "浅灰细框铝合金或低反光不锈钢体系，不用厚重黑框",
            "base": "金属槽或石材基座与挡水、防水层和地面坡向统一收口",
            "engineering": "玻璃厚度、夹层组合、最大分格、挠度、固定和五金由有资质专业单位计算深化",
        },
        "lighting": [
            {
                "type": "防潮筒灯",
                "use": "基础照明",
                "power": "建议 7–10W",
                "cct": "4000K",
                "quality": "Ra≥90，SDCM≤3",
                "protection": "汤池区灯具正面建议至少 IP65",
            },
            {
                "type": "隐藏线性灯",
                "use": "顶面和构造补光",
                "power": "按照度计算",
                "cct": "4000K",
                "quality": "24V，Ra≥90",
                "protection": "临水及易结露段建议 IP67",
            },
            {
                "type": "端景洗墙灯",
                "use": "浅浮雕端景墙",
                "power": "按配光测试",
                "cct": "4000K",
                "quality": "防眩、均匀洗墙",
                "protection": "按安装位置湿区分区复核",
            },
        ],
        "boundaries": [
            "本文件是选材、询价和施工交底草案，不是结构、防水、电气、消防、暖通、给排水或机电施工图。",
            "玻璃厚度、天花受力、灯具电气参数和湿区防护等级由有资质专业单位结合现场和当地规范深化。",
            "排水坡度、防水节点、循环水、新风、排风、除湿、消防和设备接口不在本文件构造定尺范围内。",
            "灯具数量、间距、回路、功率和配光以计算与现场样灯为准，不从效果图反推。",
        ],
    }


PRICE_REFERENCE = [
    ("池底/坐台小砖", "唯格 VIEWGRES / 佛山正规泳池砖厂家", "150–260", "80–150", "元/m²"),
    ("池壁砖", "东鹏 / 新中源同档", "120–220", "65–120", "元/m²"),
    ("池沿加工砖", "唯格 / 佛山可加工厂家", "220–380", "130–220", "元/m²"),
    ("池外及过渡区地砖", "东鹏 / 新中源、冠珠同档", "110–190", "60–110", "元/m²"),
    ("普通墙面砖", "蒙娜丽莎、东鹏 / 佛山工程砖", "130–240", "70–130", "元/m²"),
    ("端景肌理砖", "蒙娜丽莎、唯格定制 / 佛山成品肌理砖", "260–480", "140–260", "元/m²"),
    ("铝制天花系统", "佛山欧陆、欧佰同档 / 佛山工程铝扣板", "260–480", "100–180", "元/m²"),
    ("木纹铝板升级差价", "佛山工程铝板厂家 / 经济档取消", "80–160", "0", "元/m²"),
    ("防潮筒灯", "欧普 / 三雄极光、雷士同档", "90–180", "55–90", "元/套"),
    ("防水线性灯带", "欧普 / 三雄极光、广东工程厂家", "55–95", "30–55", "元/m"),
    ("端景洗墙灯", "三雄极光 / 雷士同档", "220–420", "100–220", "元/套"),
    ("24V 恒压驱动", "明纬 / 灯带同品牌配套驱动", "180–360", "90–180", "元/套"),
]


def _hex(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _set_run_font(
    run,
    *,
    name: str = FONT_NAME,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_FALLBACK)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_FALLBACK)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = _hex(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in CELL_MARGIN_DXA.items():
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def _set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def _prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def _set_table_borders(table, color: str = "C7CFCC", size: int = 4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _configure_table_geometry(table, widths: list[int]):
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        _prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths[index])
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_table_borders(table)


def _style_table_text(table, header: bool = True):
    for row_index, row in enumerate(table.rows):
        if header and row_index == 0:
            _set_repeat_header(row)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1.5)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    _set_run_font(
                        run,
                        size=8,
                        color=WHITE if header and row_index == 0 else INK,
                        bold=header and row_index == 0,
                    )
            if header and row_index == 0:
                _set_cell_shading(cell, ACCENT_DARK)
            elif row_index % 2 == 0:
                _set_cell_shading(cell, PALE_GRAY)


def _set_column_text(table, rows: list[list[str]]):
    for row_index, values in enumerate(rows):
        cells = table.rows[row_index].cells
        for col_index, value in enumerate(values):
            cells[col_index].text = str(value)


def _add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    _set_column_text(table, [headers, *rows])
    _configure_table_geometry(table, widths)
    _style_table_text(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def _add_callout(doc, title: str, text: str, *, kind: str = "note"):
    palette = {
        "note": (ACCENT_DARK, PALE_GREEN),
        "caution": (CAUTION, PALE_CAUTION),
        "risk": (RISK, PALE_RISK),
    }
    color, fill = palette[kind]
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    _set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(1)
    title_run = paragraph.add_run(f"{title}  ")
    _set_run_font(title_run, size=9, color=color, bold=True)
    _set_run_font(paragraph.add_run(text), size=8.8, color=INK)
    _configure_table_geometry(table, [CONTENT_WIDTH_DXA])
    _set_table_borders(table, color=fill, size=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def _add_body(doc, text: str, *, bold_lead: str | None = None):
    p = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        _set_run_font(p.add_run(bold_lead), size=9.5, color=INK, bold=True)
        _set_run_font(p.add_run(text[len(bold_lead) :]), size=9.5, color=INK)
    else:
        _set_run_font(p.add_run(text), size=9.5, color=INK)
    return p


def _add_bullets(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        _set_run_font(p.add_run(item), size=9.5, color=INK)


def _add_numbered(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        _set_run_font(p.add_run(item), size=9.5, color=INK)


def _add_heading(doc, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def _page_break(doc):
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _add_field(paragraph, instruction: str, placeholder: str = "1"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    _set_run_font(run, size=8.5, color=MUTED)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(17)
    section.right_margin = Mm(17)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)
    section.different_first_page_header_footer = True

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    core = doc.core_properties
    core.title = "听风径民宿汤泉 B 方案选材与构造指南"
    core.subject = "施工交底及供应商询价"
    core.author = "SkryPD"
    core.keywords = "汤泉, 瓷砖, 天花, 玻璃, 灯具, 施工交底, 供应商询价"


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_FALLBACK)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_FALLBACK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = _hex(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    body = doc.styles["Body Text"]
    body.base_style = normal
    body.paragraph_format.space_after = Pt(5)
    body.paragraph_format.line_spacing = 1.18

    title = doc.styles["Title"]
    title.font.name = FONT_NAME
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = _hex(ACCENT_DARK)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = FONT_NAME
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    subtitle.font.size = Pt(12.5)
    subtitle.font.color.rgb = _hex(MUTED)
    subtitle.paragraph_format.space_after = Pt(8)

    heading_tokens = {
        "Heading 1": (16, ACCENT_DARK, 18, 10),
        "Heading 2": (13, ACCENT, 14, 7),
        "Heading 3": (11.5, ACCENT_DARK, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _hex(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(9.5)
        style.paragraph_format.left_indent = Mm(9.5)
        style.paragraph_format.first_line_indent = Mm(-4.8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.18

    caption = doc.styles["Caption"]
    caption.font.name = FONT_NAME
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    caption.font.size = Pt(8)
    caption.font.italic = True
    caption.font.color.rgb = _hex(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)


def add_header_footer(doc: Document):
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(
        hp.add_run("听风径民宿  |  汤泉 B 方案选材与构造指南"),
        size=8,
        color=MUTED,
    )
    first_hp = section.first_page_header.paragraphs[0]
    first_hp.text = ""

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(
        fp.add_run("施工交底及供应商询价草案  |  2026-08-30  |  第 "),
        size=8,
        color=MUTED,
    )
    _add_field(fp, " PAGE ")
    _set_run_font(fp.add_run(" 页"), size=8, color=MUTED)

    first_fp = section.first_page_footer.paragraphs[0]
    first_fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(
        first_fp.add_run("施工交底及供应商询价草案  |  2026-08-30"),
        size=8,
        color=MUTED,
    )


def add_cover(doc: Document, model: dict):
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(kicker.add_run("T I N G  F E N G  J I N G"), size=9, color=ACCENT, bold=True)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(title.add_run("听风径民宿"), size=28, color=ACCENT_DARK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(4)
    _set_run_font(subtitle.add_run("汤泉 B 方案选材与构造指南"), size=20, color=INK, bold=True)
    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tag.paragraph_format.space_after = Pt(34)
    _set_run_font(tag.add_run("瓷砖  ·  天花  ·  玻璃  ·  灯具"), size=11, color=ACCENT)

    meta = [
        ["项目地点", model["project"]["location"], "文件用途", model["project"]["use"]],
        ["设计方向", model["project"]["style"], "空间条件", "全室内封闭高湿"],
        ["灯光基准", model["project"]["lighting"], "文件状态", "交底/询价草案"],
        ["版本日期", "2026-08-30", "方案状态", "B 方案已确认"],
    ]
    _add_table(doc, ["字段", "内容", "字段", "内容"], meta, [1450, 3540, 1450, 3540])
    _add_callout(
        doc,
        "文件边界",
        "本文件用于选材、询价和施工交底，不替代结构、防水、电气、消防、暖通、给排水或机电施工图；所有专业参数须结合现场和当地规范深化。",
        kind="risk",
    )
    _page_break(doc)


def add_contents(doc: Document):
    _add_heading(doc, "使用说明与目录", 1)
    _add_body(
        doc,
        "本指南把已确认的审美方向转为供应商可回复、施工方可检查的材料和构造要求。带“建议”的参数需结合送样、计算或现场复核；带“必须”的事项应作为询价和验收条件。",
    )
    rows = [
        ["01", "项目与设计基准", "明确 B 方案、色材和高湿环境"],
        ["02", "一页式选材速查", "业主、施工及供应商统一口径"],
        ["03", "瓷砖选择", "八个部位、性能、铺装与验收"],
        ["04", "天花系统", "铝蜂窝/铝扣板、检修与防腐"],
        ["05", "玻璃隔断", "安全玻璃、框架、收口和深化边界"],
        ["06", "灯具选择指南", "4000K、IP、显色、配光和检修"],
        ["07", "样板与询价", "封样流程、回复字段与送审资料"],
        ["08", "施工验收", "进场、隐蔽、安装、完工四阶段"],
        ["09", "专业边界", "排除项、待深化项和责任接口"],
        ["10", "参考附录", "采购 Excel、品牌和价格使用规则"],
    ]
    _add_table(doc, ["章节", "名称", "用途"], rows, [900, 2800, 6280])


def add_design_basis(doc: Document, model: dict):
    _add_heading(doc, "1  项目与 B 方案设计基准", 1)
    p = model["project"]
    rows = [
        ["项目", p["name"]],
        ["地点", p["location"]],
        ["区域", p["environment"]],
        ["用途", p["use"]],
        ["风格", p["style"]],
        ["色彩", "浅冷灰米白、中性灰白、浅灰白橡木；灰绿色少量点缀"],
        ["灯具 B", "小型无边/微边防潮筒灯 + 隐藏线性灯 + 端景洗墙灯；取消球形壁灯"],
        ["家具", "取消长凳，主通道保持无障碍物"],
    ]
    _add_table(doc, ["基准项", "已确认要求"], rows, [2050, 7930])
    _add_callout(
        doc,
        "颜色控制",
        "拒绝奶油黄、暖米黄、粉灰、橙红木色和明显金色纹理；所有饰面须同时在自然光与约 4000K 现场灯光下封样。",
        kind="caution",
    )
    _add_callout(
        doc,
        "高湿前提",
        "全室内封闭汤泉必须专项协调新风、排风、除湿、防结露、防霉和清洁维护；饰面选择不能代替环境控制。",
        kind="risk",
    )


def add_quick_reference(doc: Document, model: dict):
    _page_break(doc)
    _add_heading(doc, "2  一页式选材速查", 1)
    rows = [
        ["瓷砖", "浅冷灰米白哑光瓷质砖；池底/坐台 150×150，池壁 300×600，地面 600×600，墙面 600×1200", "对应型号湿态防滑资料；建议吸水率≤0.5%；锁批次"],
        ["天花", "中档 600×1200 浅灰白铝蜂窝板；经济 600×600 工程铝扣板", "面板、龙骨、吊杆、螺钉和连接件耐潮耐腐蚀"],
        ["玻璃", "通高超白透明钢化夹层安全玻璃 + 中部连续磨砂防撞带 + 防潮实体门", "厚度、分格、固定、五金和挠度由专业计算"],
        ["灯具", "小型防潮筒灯 + 24V 隐藏线性灯 + 端景洗墙灯；全部 4000K", "Ra≥90、SDCM≤3；汤区筒灯建议 IP65；临水灯带建议 IP67"],
        ["顶面点缀", "中档可用浅灰白橡木纹转印铝板，建议≤可见顶面 15%；经济档取消", "不得以天然木皮代替"],
        ["端景", "浅灰米白浅浮雕/肌理砖 + 4000K 均匀洗墙 + 单株耐湿植物", "不设水帘，不用吸水文化石"],
    ]
    _add_table(doc, ["系统", "主选方向", "下单前必须核对"], rows, [1250, 5000, 3730])
    _add_callout(
        doc,
        "禁用清单",
        "；".join(FORBIDDEN_ITEMS) + "。",
        kind="risk",
    )
    _add_body(
        doc,
        "数量口径：天花毛面积暂按现有采购模型约 196 m²，仅供概算；墙地砖、玻璃、灯槽和灯具均须以现场复尺、排版及专业深化图为准。",
        bold_lead="数量口径：",
    )


def add_tile_section(doc: Document, model: dict):
    _page_break(doc)
    _add_heading(doc, "3  瓷砖选择、性能和铺装要求", 1)
    _add_body(
        doc,
        "选砖优先级依次为：湿区安全和检测资料、易清洁与耐温湿变化、颜色与低反光、供货批次和补货能力。品牌仅作为询价起点，不替代型号级性能核对。",
    )
    rows = [
        [r["zone"], r["size"], r["selection"], r["performance"], r["execution"]]
        for r in model["tiles"]
    ]
    _add_table(
        doc,
        ["使用部位", "建议规格", "颜色/表面", "性能资料", "施工重点"],
        rows,
        [1250, 1450, 2200, 2350, 2730],
    )
    _add_heading(doc, "3.1  送样与批次", 2)
    _add_bullets(
        doc,
        [
            "报价单、检测报告、包装标签和实体样品的品牌、系列、型号、色号、规格必须一致。",
            "同一连续可见面锁定色号、尺码和生产批次；补货前核对色差，不接受只承诺“同色系”。",
            "至少 4 块组合样确认纹理频率和整体色差；池沿制作 1:1 圆角或倒圆节点样板。",
            "湿态防滑必须提供对应采购型号资料，不能用品牌目录、系列宣传或其他规格报告代替。",
        ],
    )
    _add_heading(doc, "3.2  排版、铺装与收口", 2)
    tile_execution = [
        ["排版", "先墙后地统一定位主控线；优先保证入口、池沿、端景和主要视线完整砖，设备口不得随意切出碎条。"],
        ["池体", "池底小砖配合坡向；池壁、池底、池沿尽量建立清晰对缝；坐台及入池踏步边缘圆滑、防绊。"],
        ["地面", "分区找坡至地漏或排水沟；不得以饰面厚度修补基层和防水坡向问题。"],
        ["墙面", "600×1200 墙砖先核对基层、粘结系统及必要挂贴节点；门洞、玻璃和设备口侧壁计入展开面。"],
        ["缝与角", "缝宽、填缝/美缝体系、伸缩缝和阴阳角密封由材料厂家、设计及防水专业共同确认。"],
        ["清洁", "试用运营清洁剂，确认不使砖面、填缝、密封胶和金属收口变色或失效。"],
    ]
    _add_table(doc, ["环节", "施工交底要求"], tile_execution, [1500, 8480])
    _add_callout(
        doc,
        "报价边界",
        "现有采购清单的瓷砖铺贴安装费不含找坡找平、防水、保护层、瓷砖胶、背胶、填缝/美缝和阴阳角密封材料；供应商须分别说明。",
        kind="caution",
    )


def add_ceiling_section(doc: Document, model: dict):
    _page_break(doc)
    _add_heading(doc, "4  天花系统选择、构造和检修要求", 1)
    rows = [
        [r["tier"], r["system"], r["module"], r["finish"], r["requirements"]]
        for r in model["ceilings"]
    ]
    _add_table(
        doc,
        ["档次", "系统", "建议模数", "表面", "核心要求"],
        rows,
        [1250, 1700, 1550, 1400, 4080],
    )
    _add_callout(
        doc,
        "推荐结论",
        "优先采用中档铝蜂窝板系统，整体性、平整度和灯具/风口整合能力更好；预算受限时采用工程铝扣板，但不得降级为薄型家装覆膜扣板。",
        kind="note",
    )
    components = [
        ["面板", "铝蜂窝板或工程铝扣板", "浅灰白/雾白哑光；板厚、蜂窝总厚、涂层、燃烧性能及质保写入报价。"],
        ["龙骨/吊杆", "与系统配套的高湿环境构件", "材质、镀层/涂层和防腐说明随深化图提交；不得混用来源不明辅件。"],
        ["连接件", "螺钉、挂件、收边和加固件", "适应高湿与凝露，不得使用易锈蚀普通构件污染面板。"],
        ["灯槽", "防潮铝型材 + 耐黄变柔光罩", "与面板排版、灯带散热、转角、端盖和检修方式同时深化。"],
        ["检修口", "同面材隐框可开启构造优先", "对准驱动、阀件、风机、消防和设备接口；可开启、可拆换、可维护。"],
        ["木纹点缀", "浅灰白橡木纹转印铝板", "中档建议≤可见顶面 15%，不偏黄/橙红；经济方案取消，不用天然木皮。"],
    ]
    _add_table(doc, ["组成", "建议", "询价/深化要点"], components, [1500, 2750, 5730])
    _add_heading(doc, "4.1  综合天花协调", 2)
    _add_numbered(
        doc,
        [
            "机电提供新风、排风、除湿、消防、喷淋、探测、设备和检修范围底图。",
            "吊顶厂家按现场净尺寸排板，叠加灯具、风口、喷淋和检修口，避免碎板和相互冲突。",
            "灯具与吊顶厂家共同确认开孔、背部加固、收口和可更换性；不得照搬石膏板无边灯节点。",
            "样板段确认板缝、色差、平整度、反射眩光、灯槽连续性和检修动作后再批量施工。",
        ],
    )
    _add_callout(
        doc,
        "禁用",
        "普通石膏板、天然木皮和吸水矿棉板不得用于本项目完整封闭高湿汤泉可见天花；经济方案也必须保持耐潮耐腐蚀铝制系统。",
        kind="risk",
    )


def add_glass_section(doc: Document, model: dict):
    _page_break(doc)
    _add_heading(doc, "5  通高玻璃隔断构造深化建议", 1)
    _add_callout(
        doc,
        "重要",
        "本章为构造深化建议，不是定尺施工节点。玻璃厚度、夹层组合、最大分格、挠度、固定、五金和基体承载必须由有资质专业单位计算并出具深化图。",
        kind="risk",
    )
    glass = model["glass"]
    rows = [
        ["位置与关系", glass["position"], "实体门保留，不改为全玻璃门；主要通行净空、门扇开启和逃生关系复核。"],
        ["玻璃", glass["glass_type"], glass["engineering"]],
        ["可视识别", glass["visibility"], "高度、宽度、图案和男女汤标识在深化图统一；应连续、清晰、耐清洁。"],
        ["框架", glass["frame"], "框材表面、壁厚、连接、隔离垫片和耐腐蚀能力由系统供应商提交。"],
        ["底部收口", glass["base"], "不得刺穿已完成防水而无专项修复；构造须可排除凝露和清洁积水。"],
        ["顶部/侧边", "可靠固定 + 防摆 + 允许变形密封", "与结构/墙体基体连接，预留系统允许的伸缩和密封空间。"],
        ["五金/密封", "高湿耐腐蚀五金、相容垫块及密封胶", "提交材料相容性、清洁剂耐受性、质保和更换方式。"],
    ]
    _add_table(doc, ["构造层级", "设计方向", "深化/验收重点"], rows, [1650, 3550, 4780])
    _add_heading(doc, "5.1  自上而下构造逻辑", 2)
    sequence = [
        ["顶部固定", "连接可靠基体，设置系统型材/槽口、防摆及允许变形的密封构造。"],
        ["玻璃分格", "超白透明钢化夹层安全玻璃；按计算确定厚度、夹层、最大板块和支承边。"],
        ["识别带", "中部连续磨砂防撞带，兼顾透明感、识别和适度隐私。"],
        ["门洞", "玻璃细框系统包围防潮实体门；门框、玻璃框和墙面收口连续。"],
        ["底部基座", "金属槽或石材基座与挡水、防水层、地面完成面和坡向统一。"],
        ["运维界面", "避免形成难清洁死角；密封胶、五金和玻璃边部可检查、可维修。"],
    ]
    _add_table(doc, ["顺序", "构造说明"], sequence, [1850, 8130])
    _add_heading(doc, "5.2  供应商深化图最低内容", 2)
    _add_bullets(
        doc,
        [
            "平面分格、立面分格、门洞净空、实体门开启方向和主要通行尺寸。",
            "玻璃组合、钢化/夹层标识、边部处理、孔槽限制和计算依据。",
            "顶部、侧边、底部及门框节点；所有连接基体、紧固件和密封材料标注。",
            "防水层、挡水坎/基座、完成面和地面坡向关系；穿透防水的处理方案。",
            "框料和五金材质、表面处理、耐腐蚀、质保、清洁和更换说明。",
            "磨砂防撞带、男女汤识别、可视性、无障碍和当地规范复核记录。",
        ],
    )


def add_lighting_section(doc: Document, model: dict):
    _page_break(doc)
    _add_heading(doc, "6  灯具选择指南：4000K 选择与布置", 1)
    rows = [
        [r["type"], r["use"], r["power"], r["cct"], r["quality"], r["protection"]]
        for r in model["lighting"]
    ]
    _add_table(
        doc,
        ["灯具", "用途", "功率口径", "色温", "光品质", "湿区防护"],
        rows,
        [1400, 1500, 1500, 950, 2000, 2630],
    )
    _add_callout(
        doc,
        "统一口径",
        "全部建筑照明约 4000K，不混入 3000K/3500K 暖光；送样时筒灯、灯带和洗墙灯须同场试亮，避免肉眼可见色差。",
        kind="caution",
    )
    guide = [
        ["防眩", "筒灯优先深藏光源、低亮度出光口；避免直射人眼和水面造成刺眼倒影。"],
        ["均匀度", "基础照明不追求满顶密布；通过照度计算确定数量、间距和配光，重点消除通道和池边暗区。"],
        ["洗墙", "端景洗墙灯通过样灯确定离墙距离、角度和配光，浅浮雕墙面不出现明显扇形斑或过曝。"],
        ["灯带", "24V 恒压系统按总瓦数、压降和回路分段；铝型材散热，转角、接头和末端密封可检查。"],
        ["驱动", "设在可检修、散热、防潮的位置，标注服务灯带范围和功率余量；不得密封在不可检修天花内。"],
        ["吊顶配合", "无边/微边灯在铝板上的开孔、背衬加固和收口由灯具与吊顶厂家联合深化。"],
        ["控制", "基础、线性和端景照明分回路；实际控制策略与运营、清洁和维护需求一并确认。"],
    ]
    _add_table(doc, ["主题", "选择与布置要求"], guide, [1550, 8430])
    _add_heading(doc, "6.1  样灯确认记录", 2)
    sample_rows = [
        ["色温一致性", "筒灯/灯带/洗墙灯同场 4000K，无明显偏黄、偏粉或偏绿", "[ ] 通过  [ ] 修改"],
        ["显色与色容差", "资料建议 Ra≥90、SDCM≤3，实体材料颜色自然", "[ ] 通过  [ ] 修改"],
        ["眩光和水面反射", "主要站立、坐姿和入池视点无明显刺眼光源/倒影", "[ ] 通过  [ ] 修改"],
        ["墙面均匀度", "端景无强斑、扇形或明显暗带", "[ ] 通过  [ ] 修改"],
        ["防护与检修", "IP、接头密封、驱动位置、开孔及更换方式已确认", "[ ] 通过  [ ] 修改"],
    ]
    _add_table(doc, ["测试项", "判定重点", "确认"], sample_rows, [1850, 5950, 2180])
    _add_callout(
        doc,
        "电气边界",
        "本文不含电线、线管、开关、配电箱、接地和配电改造。回路、功率、负荷、保护、等电位及湿区用电安全由电气专业按当地规范深化。",
        kind="risk",
    )


def add_sample_and_inquiry_section(doc: Document):
    _page_break(doc)
    _add_heading(doc, "7  现场样板、询价和材料送审", 1)
    _add_heading(doc, "7.1  样板确认流程", 2)
    _add_numbered(doc, SAMPLE_CHECKLIST)
    _add_heading(doc, "7.2  供应商统一回复字段", 2)
    inquiry_rows = []
    for index in range(0, len(INQUIRY_FIELDS), 4):
        row = INQUIRY_FIELDS[index : index + 4]
        inquiry_rows.append(row + [""] * (4 - len(row)))
    _add_table(doc, ["字段 1", "字段 2", "字段 3", "字段 4"], inquiry_rows, [2495] * 4)
    _add_heading(doc, "7.3  各专业最低送审资料", 2)
    documents = [
        ["瓷砖", "产品合格证；吸水率；对应型号湿态防滑资料；色号/尺码/批次；铺贴与清洁说明；池沿加工样。"],
        ["天花", "面板和涂层参数；龙骨/吊杆/连接件；防腐和燃烧性能；排板深化图；检修口及灯槽样。"],
        ["玻璃", "玻璃组合及标识；计算书/深化图；框材和五金；密封相容性；防撞带；安装与质保说明。"],
        ["灯具", "型号规格书；光通量/配光；4000K、Ra、SDCM；IP 资料；驱动参数；安装开孔和质保。"],
    ]
    _add_table(doc, ["专业", "最低资料"], documents, [1600, 8380])
    _add_callout(
        doc,
        "下单门槛",
        "书面报价、型号级检测资料、实体样板和深化图四项未齐，不得以销售口头承诺或效果图作为批量下单依据。",
        kind="risk",
    )


def add_acceptance_section(doc: Document):
    _page_break(doc)
    _add_heading(doc, "8  施工验收检查表", 1)
    _add_body(
        doc,
        "建议每个阶段由施工负责人自检，监理/设计/业主按职责复核并留存照片和签认记录。以下检查项不替代法定验收、材料复验或专业检测。",
    )
    for stage, items in ACCEPTANCE_STAGES.items():
        _add_heading(doc, f"8.{list(ACCEPTANCE_STAGES).index(stage) + 1}  {stage}", 2)
        rows = [[f"[ ] {item}", "记录：", "签认："] for item in items]
        _add_table(doc, ["检查项", "问题/整改", "责任方/日期"], rows, [6000, 2300, 1680])


def add_boundary_section(doc: Document, model: dict):
    _page_break(doc)
    _add_heading(doc, "9  专业边界与风险提示", 1)
    rows = [
        ["本指南已覆盖", "瓷砖、铝制天花、玻璃隔断设计方向和灯具选择；样板、询价、送审与饰面验收。"],
        ["不属于本指南施工定尺", "结构承载、玻璃计算、防水、找坡、保护层、给排水/循环水、暖通除湿、消防、电气和配电。"],
        ["采购清单未计入", "防水、找坡找平、保护层、瓷砖胶、填缝/美缝、给排水/循环水、电线线管、配电、新风排风除湿和消防。"],
        ["必须专业深化", "玻璃厚度/分格/固定；天花受力和机电综合；灯具照度/配光/回路/IP；排水坡度和防水节点。"],
        ["不得作为依据", "AI 效果图中的数量、尺寸、可见灯点、非几何纹理和软装细节；不得从效果图反推施工。"],
        ["当地复核", "防火、疏散、无障碍、湿区电气、安全玻璃、材料燃烧性能和施工验收均以当地主管部门及现行规范为准。"],
    ]
    _add_table(doc, ["边界类别", "说明"], rows, [2050, 7930])
    for boundary in model["boundaries"]:
        _add_callout(doc, "专业提示", boundary, kind="risk")
    _add_heading(doc, "9.1  专业接口责任清单", 2)
    responsibility = [
        ["室内/材料", "色材、排版、收口意图、样板确认和完成面验收。"],
        ["防水/给排水", "防水体系、闭水、找坡、地漏/沟、池体和循环水接口。"],
        ["暖通/除湿", "新风、排风、除湿、凝露控制、风口和检修范围。"],
        ["电气/照明", "照度、配光、回路、负荷、保护、IP、接地/等电位和控制。"],
        ["玻璃专业", "计算、分格、玻璃组合、框架、固定、五金、密封和安装质量。"],
        ["天花厂家", "排板、系统构件、防腐、灯具/风口/检修口整合和质保。"],
    ]
    _add_table(doc, ["责任专业", "最低职责"], responsibility, [1850, 8130])


def add_appendix(doc: Document):
    _page_break(doc)
    _add_heading(doc, "10  参考附录：品牌与价格的使用规则", 1)
    _add_callout(
        doc,
        "使用方法",
        "以下是 2026-08-30 采购 Excel 的广东区域询价起点，不是定标价或最终合同价。正式采购须以同一技术条件向至少三家供应商书面询价，并核对税运、安装、损耗、交期、质保和补货。",
        kind="caution",
    )
    rows = [list(row) for row in PRICE_REFERENCE]
    _add_table(
        doc,
        ["品类", "品牌/渠道方向", "中档参考", "经济参考", "单位"],
        rows,
        [1700, 3600, 1550, 1550, 1580],
    )
    _add_heading(doc, "10.1  配套采购文件", 2)
    _add_body(
        doc,
        "详细工程量、损耗、品牌价格来源、参考合价和技术验收表见《汤泉B方案瓷砖吊顶灯具采购询价清单.xlsx》。Word 指南用于统一技术口径；Excel 用于回填数量、单价和比较报价，两者应同时发送给供应商。",
    )
    _add_heading(doc, "10.2  发标前最后检查", 2)
    checklist = [
        "现场尺寸、池体和玻璃洞口已复核并形成记录。",
        "瓷砖排版、天花综合点位、玻璃深化图和灯具照度/配光已形成可询价版本。",
        "每个产品的性能资料、样板、型号和报价条件一一对应。",
        "中档与经济档采用相同技术底线比较，不以取消湿区安全或专业验收换取低价。",
        "合同明确税运、卸货、上楼、破损补发、补货色差、安装、调试、成品保护和质保责任。",
    ]
    _add_table(doc, ["发标检查项", "状态"], [[x, "[ ]"] for x in checklist], [8700, 1280])
    _add_callout(
        doc,
        "最终声明",
        "本文件为选材、询价和施工交底草案。凡涉及人身安全、结构受力、防水、电气、消防、暖通、给排水和机电系统的内容，必须由相应有资质专业单位结合现场和当地现行规范深化、审核和验收。",
        kind="risk",
    )


def build_document(output_path: str | Path = DEFAULT_OUTPUT) -> Path:
    """Build and save the approved guide, returning its absolute path."""

    output = Path(output_path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    model = build_content_model()
    doc = Document()
    configure_document(doc)
    configure_styles(doc)
    add_header_footer(doc)
    add_cover(doc, model)
    add_contents(doc)
    add_design_basis(doc, model)
    add_quick_reference(doc, model)
    add_tile_section(doc, model)
    add_ceiling_section(doc, model)
    add_glass_section(doc, model)
    add_lighting_section(doc, model)
    add_sample_and_inquiry_section(doc)
    add_acceptance_section(doc)
    add_boundary_section(doc, model)
    add_appendix(doc)
    doc.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", nargs="?", default=str(DEFAULT_OUTPUT))
    args = parser.parse_args()
    output = build_document(args.output)
    print(output)


if __name__ == "__main__":
    main()
